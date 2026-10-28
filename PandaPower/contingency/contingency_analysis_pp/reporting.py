from __future__ import annotations

from dataclasses import asdict, is_dataclass
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Iterable, Any
import math
import os

import numpy as np
import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import ContingencyResult
from .violations import Violation

# Leave margin below Excel's hard limit of 1,048,576 rows.
_MAX_DATA_ROWS_PER_SHEET = 1_000_000

_CASE_COLUMNS = [
    "case_id", "outage_type", "outage_index", "outage_name", "converged",
    "base_target_vm_pu", "post_target_vm_pu", "target_delta_vm_pu", "error",
]
_VIOLATION_COLUMNS = [
    "case_id", "kind", "element_type", "element_index", "element_name",
    "value", "limit", "base_value", "post_value", "delta_value",
    "pre_existing", "message",
]
_BRANCH_COLUMNS = [
    "case_id", "element_type", "element_index", "element_name",
    "base_loading_percent", "post_loading_percent", "delta_loading_percent",
]
_BUS_COLUMNS = [
    "case_id", "bus_index", "bus_name", "vn_kv", "base_vm_pu",
    "post_vm_pu", "delta_vm_pu",
]


def _name(net, table: str, idx: Any) -> Any:
    if table in net and idx in net[table].index and "name" in net[table].columns:
        return net[table].at[idx, "name"]
    return None


def _finite(value: Any) -> Any:
    """Replace NaN/Inf with None so openpyxl always receives valid cells."""
    if value is None:
        return None
    if isinstance(value, (float, np.floating)):
        return float(value) if math.isfinite(float(value)) else None
    if isinstance(value, np.integer):
        return int(value)
    if isinstance(value, np.bool_):
        return bool(value)
    return value


def _object_dict(obj: Any) -> dict[str, Any]:
    if is_dataclass(obj):
        return asdict(obj)
    if hasattr(obj, "__dict__"):
        return dict(vars(obj))
    slots = getattr(obj, "__slots__", ())
    return {name: getattr(obj, name) for name in slots if hasattr(obj, name)}


def _safe_frame(rows: list[dict], columns: list[str]) -> pd.DataFrame:
    frame = pd.DataFrame(rows, columns=columns)
    if frame.empty:
        return frame
    return frame.map(_finite)


def _flatten_parameters(value: Any, prefix: str = "") -> list[dict]:
    rows: list[dict] = []
    if isinstance(value, dict):
        if not value:
            rows.append({"parameter": prefix, "value": "{}", "type": "dict"})
        for key, item in value.items():
            path = f"{prefix}.{key}" if prefix else str(key)
            rows.extend(_flatten_parameters(item, path))
    elif isinstance(value, (list, tuple, set)):
        seq = sorted(value, key=str) if isinstance(value, set) else list(value)
        if not seq:
            rows.append({"parameter": prefix, "value": "[]", "type": type(value).__name__})
        for index, item in enumerate(seq):
            rows.extend(_flatten_parameters(item, f"{prefix}[{index}]"))
    else:
        rows.append({
            "parameter": prefix,
            "value": _finite(value),
            "type": "NoneType" if value is None else type(value).__name__,
        })
    return rows


def _parameter_frame(parameters: dict) -> pd.DataFrame:
    rows = _flatten_parameters(parameters)
    return _safe_frame(rows, ["parameter", "value", "type"])


def _result_rows(net, results: list[ContingencyResult]):
    cases: list[dict] = []
    buses: list[dict] = []
    branches: list[dict] = []

    for result in results:
        case_id = result.outage.case_id
        cases.append({
            "case_id": case_id,
            "outage_type": result.outage.element_type,
            "outage_index": result.outage.index,
            "outage_name": result.outage.name,
            "converged": result.converged,
            "base_target_vm_pu": result.base_v_target_pu,
            "post_target_vm_pu": result.post_v_target_pu,
            "target_delta_vm_pu": result.voltage_change_pu,
            "error": result.error,
        })

        # Keep non-convergent cases in Cases/Violations. They legitimately have
        # no post-contingency bus or branch result rows.
        all_buses = set(result.base.bus_vm_pu) | set(result.post.bus_vm_pu)
        for idx in sorted(all_buses, key=str):
            base_vm = result.base.bus_vm_pu.get(idx)
            post_vm = result.post.bus_vm_pu.get(idx)
            delta = None if base_vm is None or post_vm is None else post_vm - base_vm
            vn_kv = net.bus.at[idx, "vn_kv"] if idx in net.bus.index else None
            buses.append({
                "case_id": case_id,
                "bus_index": idx,
                "bus_name": _name(net, "bus", idx),
                "vn_kv": vn_kv,
                "base_vm_pu": base_vm,
                "post_vm_pu": post_vm,
                "delta_vm_pu": delta,
            })

        branch_sets = (
            ("line", result.base.line_loading_percent, result.post.line_loading_percent),
            ("trafo", result.base.trafo_loading_percent, result.post.trafo_loading_percent),
            ("trafo3w", result.base.trafo3w_loading_percent, result.post.trafo3w_loading_percent),
        )
        for table, base_values, post_values in branch_sets:
            for idx in sorted(set(base_values) | set(post_values), key=str):
                base_loading = base_values.get(idx)
                post_loading = post_values.get(idx)
                delta = (None if base_loading is None or post_loading is None
                         else post_loading - base_loading)
                branches.append({
                    "case_id": case_id,
                    "element_type": table,
                    "element_index": idx,
                    "element_name": _name(net, table, idx),
                    "base_loading_percent": base_loading,
                    "post_loading_percent": post_loading,
                    "delta_loading_percent": delta,
                })

    return cases, buses, branches


def _violation_rows(net, violations: list[Violation]) -> list[dict]:
    rows = []
    for violation in violations:
        row = _object_dict(violation)
        table = row.get("element_type")
        idx = row.get("element_index")
        row["element_name"] = _name(net, table, idx) if table and idx is not None else None
        rows.append(row)
    return rows


def _write_chunked(writer: pd.ExcelWriter, frame: pd.DataFrame, base_name: str) -> list[str]:
    """Write a DataFrame across sheets without exceeding Excel's row limit."""
    if frame.empty:
        frame.to_excel(writer, sheet_name=base_name, index=False)
        return [base_name]

    sheet_names = []
    for part, start in enumerate(range(0, len(frame), _MAX_DATA_ROWS_PER_SHEET), start=1):
        name = base_name if part == 1 else f"{base_name[:27]} {part}"
        frame.iloc[start:start + _MAX_DATA_ROWS_PER_SHEET].to_excel(
            writer, sheet_name=name, index=False
        )
        sheet_names.append(name)
    return sheet_names


def _format_workbook(writer: pd.ExcelWriter, sheet_names: list[str]) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for sheet_name in sheet_names:
        ws = writer.book[sheet_name]
        ws.freeze_panes = "A2"
        if ws.max_row >= 1 and ws.max_column >= 1:
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
        for column_cells in ws.iter_cols(min_row=1, max_row=min(ws.max_row, 5000)):
            width = max((len(str(cell.value)) if cell.value is not None else 0
                         for cell in column_cells), default=0)
            letter = get_column_letter(column_cells[0].column)
            ws.column_dimensions[letter].width = min(max(width + 2, 10), 45)


def write_report(path, net, results: Iterable[ContingencyResult],
                 violations: Iterable[Violation], *,
                 metadata: dict | None = None,
                 study_parameters: dict | None = None) -> Path:
    """Write a robust, atomic, multi-sheet contingency workbook.

    The workbook always contains headers, even if no contingency converged or
    no feasible connection capacity was found. Very large long-form result
    tables are automatically split across Excel worksheets.
    """
    path = Path(path).expanduser()
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.suffix.lower() != ".xlsx":
        path = path.with_suffix(".xlsx")

    results = list(results)
    violations = list(violations)
    case_rows, bus_rows, branch_rows = _result_rows(net, results)
    violation_rows = _violation_rows(net, violations)

    meta = dict(metadata or {})
    meta.update({
        "case_count_written": len(results),
        "violation_count_written": len(violations),
        "bus_result_rows": len(bus_rows),
        "branch_result_rows": len(branch_rows),
        "report_status": (
            "OK" if results else
            "NO CASE RESULTS: inspect capacity search/base N-1 feasibility"
        ),
    })

    study_df = _safe_frame([meta], list(meta.keys()))
    cases_df = _safe_frame(case_rows, _CASE_COLUMNS)
    violations_df = _safe_frame(violation_rows, _VIOLATION_COLUMNS)
    branches_df = _safe_frame(branch_rows, _BRANCH_COLUMNS)
    buses_df = _safe_frame(bus_rows, _BUS_COLUMNS)
    parameter_df = _parameter_frame(study_parameters or {})

    # Write to a temporary file in the destination directory, then atomically
    # replace the target. A failed export cannot leave a half-written workbook.
    with NamedTemporaryFile(prefix=f".{path.stem}_", suffix=".xlsx",
                            dir=path.parent, delete=False) as tmp:
        temp_path = Path(tmp.name)

    try:
        with pd.ExcelWriter(temp_path, engine="openpyxl") as writer:
            study_df.to_excel(writer, sheet_name="Study", index=False)
            parameter_df.to_excel(writer, sheet_name="Parameters", index=False)
            cases_df.to_excel(writer, sheet_name="Cases", index=False)
            violations_df.to_excel(writer, sheet_name="Violations", index=False)
            written = ["Study", "Parameters", "Cases", "Violations"]
            written += _write_chunked(writer, branches_df, "Branch results")
            written += _write_chunked(writer, buses_df, "Bus results")
            _format_workbook(writer, written)
        os.replace(temp_path, path)
    except Exception:
        temp_path.unlink(missing_ok=True)
        raise

    return path
