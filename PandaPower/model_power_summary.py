"""Print/export installed and operating powers, reporting batteries separately.

Dependencies: pandapower, pandas, numpy; openpyxl for Excel export.

    tables = print_model_powers('model.json', output_dir='reports')
    tables = print_model_powers('model.json', 'reports', run_powerflow=True)

Totals are EXCLUSIVE: load, sgen, gen exclude identified batteries. The battery
row combines identified batteries from all three tables, with generation signs:
positive P = discharging, negative P = charging; positive Q = injection.
No network elements or original JSON file are modified by reporting.

Battery identification uses explicit technology/type tags, including the
builder's BATTERYSTRG tag. No name matching, battery_mode, or battery_state is
used for detection. Customize battery_technologies for additional exact labels.
State is derived from operating P, not potentially stale battery_state metadata.
"idle" means zero active exchange within battery_power_tolerance_mw; Q may be
nonzero. In-service elements with unavailable P have state 'unknown'.

Battery capacity uses explicit installed/rated columns (including
rated_power_mw), then max_p_mw; never the partially dispatched p_mw or sn_mva
as an assumed MW rating. Installed Q limits on battery loads are converted to
generation signs: Qmin_new=-Qmax_old, Qmax_new=-Qmin_old. S_inst sums sn_mva
or explicit installed MVA ratings. Installed ratings include out-of-service
rows, without scaling. Missing/partial sums are N/A, never silently zero.

Operating mode 'setpoints': scaled load/sgen P/Q and gen P. Gen Q/slack P
come from saved converged results or remain N/A. This is not a solved state:
ZIP loads assume nominal voltage and disconnected islands are not detected.
Mode 'results': use saved res_* values without applying scaling again, requiring
net.converged. run_powerflow=True recalculates and selects 'results'. Saved
convergence alone cannot guarantee freshness. Out-of-service elements/buses
contribute zero. S_op = hypot(sum(P),sum(Q)), not sum of element magnitudes.

Battery breakdowns include separate P_charge/P_discharge [MW] (nonnegative)
so simultaneously charging/discharging devices do not disappear in net totals.
Rows are counted once per JSON element; separate duplicate representations of
the same physical battery cannot be deduplicated without a reliable asset ID.
Dedicated net.storage, ext_grid and asymmetric elements remain out of scope.

Output: console plus optional <JSON stem>_power_summary.xlsx in output_dir,
created as needed. Existing same-name workbooks are replaced. Every table,
including empty tables, gets its own sheet. Values are numeric snapshots at
full precision; decimals controls display. Returns the dictionary of tables.
"""
from __future__ import annotations

import argparse
from datetime import datetime
import os
from pathlib import Path
import re
import sys
import tempfile
from typing import Mapping, Sequence, TextIO
import unicodedata

import numpy as np
import pandas as pd
import pandapower as pp

_TECH_COLUMNS = ('technology', 'Technológia', 'tech')
_SOURCE_COLUMNS = ('source_file', 'source_filename', 'source_path', 'filename', 'Source', 'source')
_BATTERY_TECHNOLOGIES = ('BATTERYSTRG', 'BATTERY', 'BATTERY_STORAGE', 'BESS', 'AKKUMULÁTOR', 'AKKUMULÁTOROS ENERGIATÁROLÓ')
_RATING_COLUMNS = {
    'p': ('installed_p_mw', 'p_installed_mw', 'rated_power_mw', 'rated_p_mw', 'p_rated_mw', 'p_nom_mw', 'pn_mw', 'max_p_mw'),
    'q_min': ('installed_min_q_mvar', 'min_q_mvar'),
    'q_max': ('installed_max_q_mvar', 'installed_q_mvar', 'q_installed_mvar', 'max_q_mvar'),
    's': ('installed_s_mva', 's_installed_mva', 'sn_mva'),
}
_HEADERS = {
    'n': 'N', 'n_on': 'N_on', 'p_inst_mw': 'P_inst [MW]',
    'q_min_mvar': 'Q_min [MVAr]', 'q_max_mvar': 'Q_max [MVAr]',
    's_inst_mva': 'S_inst [MVA]', 'p_op_mw': 'P_op [MW]',
    'q_op_mvar': 'Q_op [MVAr]', 's_op_mva': 'S_op [MVA]',
    'p_charge_mw': 'P_charge [MW]', 'p_discharge_mw': 'P_discharge [MW]',
}
_BASE_COLUMNS = tuple(_HEADERS)[:-2]


def _key(value):
    text = unicodedata.normalize('NFKD', str(value).casefold())
    return ''.join(char for char in text if char.isalnum())


def _column_names(df, names):
    lookup = {_key(c): c for c in df.columns}
    return list(dict.fromkeys(lookup[_key(n)] for n in names if _key(n) in lookup))


def _number(df, names):
    result = pd.Series(np.nan, index=df.index, dtype=float)
    for column in _column_names(df, names):
        values = pd.to_numeric(df[column], errors='coerce').astype(float)
        result = result.fillna(values.where(np.isfinite(values)))
    return result


def _label(value):
    if isinstance(value, (list, tuple, set, np.ndarray)):
        return ' | '.join(sorted({text for item in value if (text := _label(item))}))
    if isinstance(value, dict):
        return '; '.join(f'{k}={_label(v)}' for k, v in sorted(value.items(), key=lambda item: str(item[0])))
    return '' if value is None or bool(pd.isna(value)) else str(value).strip()


def _labels(df, names):
    result = pd.Series('', index=df.index, dtype=object)
    for column in _column_names(df, names):
        result = result.mask(result.eq(''), df[column].map(_label))
    return result.mask(result.eq(''), '<unknown>')


def _in_service(df):
    return df['in_service'].fillna(False).eq(True) if 'in_service' in df else pd.Series(True, index=df.index)


def _battery_mask(df, technology_columns, aliases):
    def matches(value):
        if isinstance(value, (list, tuple, set, np.ndarray)):
            return any(matches(item) for item in value)
        return any(_key(token) in aliases for token in re.split(r'[|;]', _label(value)))
    mask = pd.Series(False, index=df.index)
    for column in _column_names(df, (*technology_columns, 'type')):
        mask |= df[column].map(matches).astype(bool)
    return mask


def _strict_sum(values):
    return float(values.sum()) if values.notna().all() else float('nan')


def _aggregate(frame, battery=False):
    powers = {name: _strict_sum(frame[name]) for name in (
        'p_inst_mw', 'q_min_mvar', 'q_max_mvar', 's_inst_mva', 'p_op_mw', 'q_op_mvar')}
    result = {'n': len(frame), 'n_on': int(frame['on'].sum()), **powers,
              's_op_mva': float(np.hypot(powers['p_op_mw'], powers['q_op_mvar']))}
    if battery:
        result['p_charge_mw'] = _strict_sum((-frame['p_op_mw']).clip(lower=0))
        result['p_discharge_mw'] = _strict_sum(frame['p_op_mw'].clip(lower=0))
    return result


def _group(frame, columns, battery=False):
    rows = []
    for labels, subset in frame.groupby(columns, sort=True, dropna=False):
        if not isinstance(labels, tuple):
            labels = (labels,)
        rows.append({**dict(zip(columns, labels)), **_aggregate(subset, battery)})
    fields = list(_HEADERS) if battery else list(_BASE_COLUMNS)
    return pd.DataFrame(rows, columns=columns + fields).set_index(columns)


def _excel_path(json_file, output_dir, excel_filename):
    filename = excel_filename if excel_filename is not None else f'{json_file.stem}_power_summary.xlsx'
    if (not isinstance(filename, str) or not filename or '/' in filename or '\\' in filename
            or Path(filename).suffix.lower() != '.xlsx'):
        raise ValueError('excel_filename must be an .xlsx filename, not a path')
    return Path(output_dir).expanduser() / filename


def _save_tables_excel(tables, destination, *, json_file, operational, run_powerflow, notes, decimals):
    """Export precomputed tables as a numeric snapshot, without modifying them."""
    try:
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.hyperlink import Hyperlink
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError as exc:
        raise ImportError('Excel export requires: pip install openpyxl') from exc
    destination.parent.mkdir(parents=True, exist_ok=True)
    body_font = Font(name='Calibri', size=11)
    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    heading_font = Font(name='Calibri', size=11, bold=True)
    link_font = Font(name='Calibri', size=11, color='0563C1', underline='single')
    header_fill = PatternFill('solid', fgColor='24476A')
    left = Alignment(horizontal='left', vertical='center', wrap_text=True)
    right = Alignment(horizontal='right', vertical='center')
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    fmt = '#,##0' + ('.' + '0' * decimals if decimals else '')
    fd, temp_name = tempfile.mkstemp(prefix='.power_summary_', suffix='.xlsx', dir=destination.parent)
    os.close(fd)
    temporary = Path(temp_name)
    try:
        with pd.ExcelWriter(temporary, engine='openpyxl') as writer:
            for number, (name, table) in enumerate(tables.items(), 1):
                flat = table.rename(columns=_HEADERS).reset_index()
                flat.to_excel(writer, sheet_name=name, index=False, na_rep='N/A')
                ws = writer.sheets[name]
                labels = table.index.nlevels
                ws.sheet_view.showGridLines = False
                ws.freeze_panes = f'{get_column_letter(labels + 1)}2'
                ws.row_dimensions[1].height = 32
                for row in ws:
                    for cell in row:
                        cell.font = body_font
                        cell.alignment = left if cell.column <= labels else right
                        if isinstance(cell.value, str):
                            cell.data_type = 's'
                        if cell.row == 1:
                            cell.font, cell.fill, cell.alignment = header_font, header_fill, center
                        elif cell.column > labels and isinstance(cell.value, (int, float)):
                            cell.number_format = '#,##0' if flat.columns[cell.column - 1] in {'N', 'N_on'} else fmt
                    if row[0].row > 1:
                        ws.row_dimensions[row[0].row].height = 22
                for col, name_col in enumerate(flat.columns, 1):
                    width = max([len(str(name_col))] + [len(str(v)) for v in flat.iloc[:, col-1]]) + 2
                    ws.column_dimensions[get_column_letter(col)].width = min(55, max(16, width))
                if len(flat):
                    tab = Table(displayName=f'PowerSummary{number}', ref=ws.dimensions)
                    tab.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=False)
                    ws.add_table(tab)
            ws = writer.sheets['totals']
            last_col, row = ws.max_column, ws.max_row + 3
            context = [
                ('Source JSON', str(json_file.resolve())),
                ('Generated', datetime.now().astimezone().isoformat(timespec='seconds')),
                ('Operational mode', operational),
                ('Result freshness', 'Recalculated for this export' if run_powerflow else 'Not recalculated; saved results may be stale'),
                ('Exclusive totals', 'load/sgen/gen exclude identified batteries. battery combines them once per JSON element.'),
                ('Battery signs', 'Positive P = discharge, negative P = charge; positive Q = injection. Battery load P/Q and Q limits are sign-converted.'),
                ('Battery state', 'Derived from reported P, not battery_state metadata. Idle means zero active exchange within tolerance, not necessarily Q=0.'),
                ('Installed ratings', 'All equipment, no scaling. Explicit installed/rated columns (including rated_power_mw), then max_p_mw; S from MVA ratings.'),
                ('Reactive limits', 'Independent Q_min/Q_max limits in MVAr, not a simultaneous operating point. Missing limits remain N/A.'),
                ('Apparent power', 'S_inst sums ratings. S_op = hypot(sum(P_op),sum(Q_op)); not sum of device magnitudes.'),
                ('Battery gross flows', 'P_charge/P_discharge are nonnegative separate MW magnitudes, preventing cancellation of simultaneous charge/discharge.'),
                ('Operating scope', 'In-service elements/buses. Setpoints assume nominal voltage for ZIP loads and do not detect islands.'),
                ('Other signs', 'Non-battery load positive P = consumption; non-battery gen/sgen positive P = generation.'),
                ('Missing values', 'Literal text N/A is missing/incomplete information, never an Excel error or zero.'),
                ('Snapshot', 'Full-precision numeric export, display rounded; no live model or Excel power-flow calculations.'),
                ('Scope', 'load, sgen, gen. Excludes net.storage, ext_grid and asymmetric elements. No asset-level deduplication.'),
            ]
            for label, text in context + [('Note', note) for note in dict.fromkeys(notes)]:
                cell = ws.cell(row=row, column=1, value=label)
                cell.font, cell.alignment = heading_font, left
                ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=last_col)
                cell = ws.cell(row=row, column=2, value=text)
                cell.data_type, cell.font, cell.alignment = 's', body_font, left
                ws.row_dimensions[row].height = 34
                row += 1
            row += 1
            ws.cell(row=row, column=1, value='Worksheets').font = heading_font
            for name in tables:
                row += 1
                ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=last_col)
                cell = ws.cell(row=row, column=2, value=name)
                cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{name}'!A1")
                cell.font = link_font
                ws.row_dimensions[row].height = 20
        os.replace(temporary, destination)
    finally:
        temporary.unlink(missing_ok=True)



def collect_model_power_tables(
    net,
    *,
    operational: str = 'setpoints',
    technology_columns: Sequence[str] = _TECH_COLUMNS,
    source_columns: Sequence[str] = _SOURCE_COLUMNS,
    installed_columns: Mapping[str, Mapping[str, str | Sequence[str]]] | None = None,
    battery_technologies: Sequence[str] = _BATTERY_TECHNOLOGIES,
    battery_power_tolerance_mw: float = 1e-9,
) -> dict[str, pd.DataFrame]:
    """Collect the same power-summary tables as :func:`print_model_powers`.

    Unlike ``print_model_powers``, this function operates on an existing
    pandapower network and performs no printing, JSON loading, power flow or
    Excel export. It is intended for embedding the exact same summary data in
    other reports. The supplied network is not modified.
    """
    if operational not in {'setpoints', 'results'}:
        raise ValueError("operational must be 'setpoints' or 'results'")
    if not np.isfinite(battery_power_tolerance_mw) or battery_power_tolerance_mw < 0:
        raise ValueError('battery_power_tolerance_mw must be finite and nonnegative')
    if isinstance(battery_technologies, str):
        raise TypeError('battery_technologies must be a sequence of labels, not one string')
    for element, fields in (installed_columns or {}).items():
        if element not in {'load', 'sgen', 'gen'} or set(fields) - set(_RATING_COLUMNS):
            raise ValueError(f'Invalid installed_columns entry for {element!r}')
    aliases = {_key(label) for label in battery_technologies if _key(label)}
    converged = bool(net.get('converged', False))
    if operational == 'results' and not converged:
        raise ValueError("No converged results. Run a power flow or use operational='setpoints'.")
    bus_on = _in_service(net.get('bus', pd.DataFrame()))
    frames, batteries = {}, []
    for element in ('load', 'sgen', 'gen'):
        df = net.get(element, pd.DataFrame())
        data = pd.DataFrame(index=df.index)
        is_battery = _battery_mask(df, technology_columns, aliases)
        data['technology'] = _labels(df, technology_columns)
        data.loc[is_battery & data['technology'].eq('<unknown>'), 'technology'] = _labels(df, ('type',))[is_battery]
        data['source'] = _labels(df, source_columns)
        data['element'] = element
        on = _in_service(df)
        if len(df) and 'bus' in df:
            on &= df['bus'].map(bus_on).fillna(False).astype(bool)
        data['on'] = on
        overrides = (installed_columns or {}).get(element, {})
        for metric, output in (('p', 'p_inst_mw'), ('q_min', 'q_min_mvar'),
                               ('q_max', 'q_max_mvar'), ('s', 's_inst_mva')):
            names = overrides.get(metric, _RATING_COLUMNS[metric])
            data[output] = _number(df, (names,) if isinstance(names, str) else names)
            if metric == 's':
                data[output] = data[output].where(data[output] >= 0)
        result = net.get(f'res_{element}', pd.DataFrame()).reindex(df.index)
        if operational == 'results':
            p_op = _number(result, ('p_mw',))
            q_op = _number(result, ('q_mvar',))
        else:
            scale = _number(df, ('scaling',)) if 'scaling' in df else pd.Series(1., index=df.index)
            p_op = _number(df, ('p_mw',)) * scale
            if element == 'gen':
                q_op = _number(result, ('q_mvar',)) if converged else pd.Series(np.nan, index=df.index)
                slack = df['slack'].fillna(False).eq(True) if 'slack' in df else pd.Series(False, index=df.index)
                slack_p = _number(result, ('p_mw',)) if converged else pd.Series(np.nan, index=df.index)
                p_op = p_op.where(~slack, slack_p)
            else:
                q_op = _number(df, ('q_mvar',)) * scale
        data['p_op_mw'] = p_op.where(on, 0.)
        data['q_op_mvar'] = q_op.where(on, 0.)
        frames[element] = data.loc[~is_battery].copy()
        battery_part = data.loc[is_battery].copy()
        if element == 'load':
            battery_part[['p_op_mw', 'q_op_mvar']] *= -1
            q_min, q_max = battery_part['q_min_mvar'].copy(), battery_part['q_max_mvar'].copy()
            battery_part['q_min_mvar'], battery_part['q_max_mvar'] = -q_max, -q_min
        batteries.append(battery_part)
    battery = pd.concat(batteries, ignore_index=True)
    battery['state'] = 'unknown'
    battery.loc[battery['p_op_mw'].lt(-battery_power_tolerance_mw), 'state'] = 'charging'
    battery.loc[battery['p_op_mw'].gt(battery_power_tolerance_mw), 'state'] = 'discharging'
    battery.loc[battery['p_op_mw'].abs().le(battery_power_tolerance_mw), 'state'] = 'idle'
    battery.loc[~battery['on'], 'state'] = 'out_of_service'
    totals = [_aggregate(frames[element]) for element in ('load', 'sgen', 'gen')] + [_aggregate(battery)]
    tables = {
        'totals': pd.DataFrame(
            totals, index=pd.Index(('load', 'sgen', 'gen', 'battery'), name='element')
        )
    }
    for element in ('sgen', 'gen'):
        for suffix, keys in (
            ('technology', ['technology']),
            ('source', ['source']),
            ('technology_source', ['technology', 'source']),
        ):
            tables[f'{element}_by_{suffix}'] = _group(frames[element], keys)
    for suffix, keys in (
        ('technology', ['technology']),
        ('source', ['source']),
        ('technology_source', ['technology', 'source']),
        ('state', ['state']),
        ('element', ['element']),
    ):
        tables[f'battery_by_{suffix}'] = _group(battery, keys, battery=True)
    return tables

def print_model_powers(
    json_file: str | Path,
    output_dir: str | Path | None = None,
    *,
    excel_filename: str | None = None,
    operational: str = 'setpoints',
    run_powerflow: bool = False,
    runpp_kwargs: Mapping | None = None,
    technology_columns: Sequence[str] = _TECH_COLUMNS,
    source_columns: Sequence[str] = _SOURCE_COLUMNS,
    installed_columns: Mapping[str, Mapping[str, str | Sequence[str]]] | None = None,
    battery_technologies: Sequence[str] = _BATTERY_TECHNOLOGIES,
    battery_power_tolerance_mw: float = 1e-9,
    decimals: int = 3,
    stream: TextIO | None = None,
) -> dict[str, pd.DataFrame]:
    """Return exclusive load/sgen/gen/battery totals and their breakdown tables.

    output_dir=None preserves print-only behavior. Otherwise create the output
    directory and save all tables as <JSON stem>_power_summary.xlsx (overwriting).
    installed_columns overrides are keyed by ORIGINAL element table, even for
    batteries: load/sgen/gen -> p, q_min, q_max, s -> column or ordered columns.
    battery_technologies replaces the exact tag aliases (technology or type).
    """
    if operational not in {'setpoints', 'results'}:
        raise ValueError("operational must be 'setpoints' or 'results'")
    if not isinstance(decimals, int) or decimals < 0:
        raise ValueError('decimals must be a nonnegative integer')
    if not np.isfinite(battery_power_tolerance_mw) or battery_power_tolerance_mw < 0:
        raise ValueError('battery_power_tolerance_mw must be finite and nonnegative')
    if isinstance(battery_technologies, str):
        raise TypeError('battery_technologies must be a sequence of labels, not one string')
    aliases = {_key(label) for label in battery_technologies if _key(label)}
    if runpp_kwargs and not run_powerflow:
        raise ValueError('runpp_kwargs requires run_powerflow=True')
    if excel_filename is not None and output_dir is None:
        raise ValueError('excel_filename requires output_dir')
    for element, fields in (installed_columns or {}).items():
        if element not in {'load', 'sgen', 'gen'} or set(fields) - set(_RATING_COLUMNS):
            raise ValueError(f'Invalid installed_columns entry for {element!r}')
    path = Path(json_file).expanduser()
    if not path.is_file():
        raise FileNotFoundError(path)
    destination = _excel_path(path, output_dir, excel_filename) if output_dir is not None else None
    net = pp.from_json(str(path))
    if run_powerflow:
        pp.runpp(net, **dict(runpp_kwargs or {}))
        operational = 'results'
    converged = bool(net.get('converged', False))
    if operational == 'results' and not converged:
        raise ValueError("No converged results. Use run_powerflow=True or operational='setpoints'.")
    bus_on = _in_service(net.get('bus', pd.DataFrame()))
    frames, batteries, notes = {}, [], []
    if operational == 'setpoints':
        notes.append('Setpoints are not a solved state; ZIP loads assume nominal voltage, and disconnected islands are not detected.')
    if converged and not run_powerflow:
        notes.append('Saved convergence does not prove result freshness; use run_powerflow=True to recalculate.')
    for element in ('load', 'sgen', 'gen'):
        df = net.get(element, pd.DataFrame())
        data = pd.DataFrame(index=df.index)
        is_battery = _battery_mask(df, technology_columns, aliases)
        data['technology'] = _labels(df, technology_columns)
        data.loc[is_battery & data['technology'].eq('<unknown>'), 'technology'] = _labels(df, ('type',))[is_battery]
        data['source'] = _labels(df, source_columns)
        data['element'] = element
        on = _in_service(df)
        if len(df) and 'bus' in df:
            on &= df['bus'].map(bus_on).fillna(False).astype(bool)
        data['on'] = on
        overrides = (installed_columns or {}).get(element, {})
        for metric, output in (('p', 'p_inst_mw'), ('q_min', 'q_min_mvar'), ('q_max', 'q_max_mvar'), ('s', 's_inst_mva')):
            names = overrides.get(metric, _RATING_COLUMNS[metric])
            data[output] = _number(df, (names,) if isinstance(names, str) else names)
            if metric == 's':
                data[output] = data[output].where(data[output] >= 0)
            if data[output].notna().sum() < len(df):
                notes.append(f'{element}: {output} available for {data[output].notna().sum()}/{len(df)} rows; incomplete group sums are N/A.')
        result = net.get(f'res_{element}', pd.DataFrame()).reindex(df.index)
        if operational == 'results':
            p, q = _number(result, ('p_mw',)), _number(result, ('q_mvar',))
        else:
            scale = _number(df, ('scaling',)) if 'scaling' in df else pd.Series(1., index=df.index)
            p = _number(df, ('p_mw',)) * scale
            if element == 'gen':
                q = _number(result, ('q_mvar',)) if converged else pd.Series(np.nan, index=df.index)
                slack = df['slack'].fillna(False).eq(True) if 'slack' in df else pd.Series(False, index=df.index)
                slack_p = _number(result, ('p_mw',)) if converged else pd.Series(np.nan, index=df.index)
                p = p.where(~slack, slack_p)
                if len(df) and converged:
                    notes.append('gen Q and slack-gen P use saved results; other P uses scaled setpoints.')
            else:
                q = _number(df, ('q_mvar',)) * scale
        data['p_op_mw'], data['q_op_mvar'] = p.where(on, 0.), q.where(on, 0.)
        for column in ('p_op_mw', 'q_op_mvar'):
            if data[column].isna().any():
                notes.append(f'{element}: {column} missing for {data[column].isna().sum()} in-service rows; affected sums are N/A.')
        frames[element] = data.loc[~is_battery].copy()
        bat = data.loc[is_battery].copy()
        if element == 'load':
            bat[['p_op_mw', 'q_op_mvar']] *= -1
            q_min, q_max = bat['q_min_mvar'].copy(), bat['q_max_mvar'].copy()
            bat['q_min_mvar'], bat['q_max_mvar'] = -q_max, -q_min
        batteries.append(bat)
        if len(bat):
            notes.append(f'Extracted {len(bat)} battery rows from {element}; excluded from non-battery totals and breakdowns.')
    battery = pd.concat(batteries, ignore_index=True)
    battery['state'] = 'unknown'
    battery.loc[battery['p_op_mw'].lt(-battery_power_tolerance_mw), 'state'] = 'charging'
    battery.loc[battery['p_op_mw'].gt(battery_power_tolerance_mw), 'state'] = 'discharging'
    battery.loc[battery['p_op_mw'].abs().le(battery_power_tolerance_mw), 'state'] = 'idle'
    battery.loc[~battery['on'], 'state'] = 'out_of_service'
    notes.append(f'Battery idle threshold: abs(P_op) <= {battery_power_tolerance_mw:g} MW; Q may be nonzero.')
    if not len(battery):
        notes.append('No batteries matched the technology/type aliases; provide battery_technologies if your labels differ.')
    for field in ('technology', 'source'):
        if battery[field].eq('<unknown>').any():
            notes.append(f'Battery {field} unknown for {battery[field].eq("<unknown>").sum()} rows; check metadata column aliases.')
    if len(net.get('storage', [])):
        notes.append('net.storage is present but outside this function scope; it is not included in the battery tables.')
    totals = [_aggregate(frames[element]) for element in ('load', 'sgen', 'gen')] + [_aggregate(battery)]
    tables = {'totals': pd.DataFrame(totals, index=pd.Index(('load', 'sgen', 'gen', 'battery'), name='element'))}
    for element in ('sgen', 'gen'):
        for suffix, keys in (('technology', ['technology']), ('source', ['source']), ('technology_source', ['technology', 'source'])):
            tables[f'{element}_by_{suffix}'] = _group(frames[element], keys)
    for suffix, keys in (('technology', ['technology']), ('source', ['source']),
                         ('technology_source', ['technology', 'source']), ('state', ['state']), ('element', ['element'])):
        tables[f'battery_by_{suffix}'] = _group(battery, keys, battery=True)
    out = sys.stdout if stream is None else stream
    print(f'\nMODEL: {path}\nOperational mode: {operational}', file=out)
    print('EXCLUSIVE totals: load/sgen/gen EXCLUDE batteries. Installed ratings include out-of-service equipment, no scaling.', file=out)
    print('Battery P: positive = discharge, negative = charge. Battery Q: positive = injection.', file=out)
    print('Other P signs: load positive = consumption; gen/sgen positive = generation.', file=out)
    print('S_inst = sum of ratings; S_op = hypot(sum P, sum Q). Q_min/Q_max are independent installed limits.', file=out)
    for title, table in tables.items():
        if not table.empty:
            print(f'\n{title.upper().replace("_", " ")}', file=out)
            print(table.rename(columns=_HEADERS).to_string(na_rep='N/A', float_format=lambda v: f'{v:,.{decimals}f}'), file=out)
    print('\nNOTES', file=out)
    for note in dict.fromkeys(notes):
        print(f'- {note}', file=out)
    if destination is not None:
        _save_tables_excel(tables, destination, json_file=path, operational=operational,
                           run_powerflow=run_powerflow, notes=notes, decimals=decimals)
        print(f'\nExcel saved: {destination.resolve()}', file=out)
    return tables


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('json_file', type=Path)
    parser.add_argument('--output-dir', '-o', type=Path)
    parser.add_argument('--excel-filename')
    parser.add_argument('--operational', choices=('setpoints', 'results'), default='setpoints')
    parser.add_argument('--run-powerflow', action='store_true')
    parser.add_argument('--decimals', type=int, default=3)
    args = parser.parse_args()
    print_model_powers(args.json_file, args.output_dir, excel_filename=args.excel_filename,
                       operational=args.operational, run_powerflow=args.run_powerflow, decimals=args.decimals)
